# -*- coding: utf-8 -*-
"""
Rebuild bilingual Excel from ORIGINAL file + questions.json Chinese.
Keeps original structure; inserts 中文 column after English.
"""
from copy import copy
from pathlib import Path
import json
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
ORIG = ROOT / "Impact 2027-data collection support v1.xlsx"
OUT = ROOT / "Impact 2027-data collection support v1_雙語提問.xlsx"
QPATH = Path(__file__).resolve().parent / "data" / "questions.json"


def unmerge_all(ws):
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def main():
    data = json.loads(QPATH.read_text(encoding="utf-8"))
    by_ref = {str(q["ref"]): q for q in data["questions"]}
    for m in data.get("metrics", []):
        by_ref[str(m["ref"])] = m
    # option lookup by en under parent not needed for sheet rebuild of col C only

    opt_zh = {}
    for q in data["questions"]:
        for opt in q.get("options") or []:
            if opt.get("en") and opt.get("zh"):
                opt_zh[str(opt["en"]).strip()] = opt["zh"]

    wb = load_workbook(ORIG)
    for name in wb.sheetnames:
        if not name.strip().startswith("SDG"):
            continue
        ws = wb[name]
        unmerge_all(ws)
        ws.insert_cols(4)  # new D = 中文
        # title ZH
        sdg = re.match(r"SDG\s*(\d+)", name.strip(), re.I)
        if sdg:
            meta = data["sdgs"].get(sdg.group(1), {})
            ws["D1"] = meta.get("zh") or ""
        ws["C2"] = "English (Question)"
        ws["D2"] = "中文（提問）"
        # shift headers already by insert; fix E-H from old D-G labels if row2 had them
        # Original: D Value, E Yes/No, F Evidence, G Public -> now E F G H
        ws["E2"] = "Value\n(for continuous data)"
        ws["F2"] = "Yes/No"
        ws["G2"] = "Evidence1"
        ws["H2"] = "Public (Yes/No)"

        wrap = Alignment(wrap_text=True, vertical="top")
        gray = PatternFill("solid", fgColor="D9D9D9")

        for row in range(3, ws.max_row + 1):
            typ = ws.cell(row, 1).value
            ref = ws.cell(row, 2).value
            en = ws.cell(row, 3).value
            typ_s = str(typ).strip() if typ else ""
            ref_s = str(ref).strip() if ref is not None else ""
            en_s = str(en).strip() if en else ""

            zh = ""
            if ref_s and ref_s in by_ref:
                zh = by_ref[ref_s].get("zh") or ""
            elif en_s and en_s in opt_zh:
                zh = opt_zh[en_s]

            ws.cell(row, 4).value = zh or None
            ws.cell(row, 3).alignment = wrap
            ws.cell(row, 4).alignment = wrap

            # gray Metric rows (不用填)
            if typ_s == "Metric":
                for col in range(1, 9):
                    cell = ws.cell(row, col)
                    if not cell.fill or cell.fill.fill_type is None or (
                        getattr(cell.fill.fgColor, "rgb", None) in (None, "00000000")
                    ):
                        cell.fill = gray

        ws.column_dimensions["C"].width = 55
        ws.column_dimensions["D"].width = 55

    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
