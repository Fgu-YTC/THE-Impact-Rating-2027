# -*- coding: utf-8 -*-
"""Audit original Excel structure: Metric vs Indicator vs sub-rows."""
from pathlib import Path
from openpyxl import load_workbook
import json

ROOT = Path(__file__).resolve().parent.parent
ORIG = ROOT / "Impact 2027-data collection support v1.xlsx"
OUT = Path(__file__).resolve().parent / "data" / "_audit_original.json"

wb = load_workbook(ORIG)
report = {}

for name in wb.sheetnames:
    if not name.strip().startswith("SDG"):
        continue
    ws = wb[name]
    rows = []
    for row in range(3, ws.max_row + 1):
        typ = ws.cell(row, 1).value
        ref = ws.cell(row, 2).value
        label = ws.cell(row, 3).value
        if typ is None and ref is None and label is None:
            continue
        # answer-area fills D-G
        fills = []
        for col in range(4, 8):
            f = ws.cell(row, col).fill
            if f and f.fill_type:
                fg = f.fgColor
                if fg and fg.type == "theme":
                    fills.append(f"t{fg.theme}")
                elif fg and fg.type == "rgb":
                    fills.append(str(fg.rgb))
                else:
                    fills.append(f.fill_type)
            else:
                fills.append("-")
        # Heuristic: Metric rows often have more grayed answer cols
        gray_count = sum(1 for x in fills if x != "-")
        rows.append(
            {
                "row": row,
                "type": typ,
                "ref": str(ref) if ref is not None else None,
                "label": str(label).strip() if label else None,
                "answerFills": fills,
                "grayedCols": gray_count,
            }
        )
    report[name.strip()] = {
        "title": ws["A1"].value,
        "metrics": sum(1 for r in rows if r["type"] == "Metric"),
        "indicators": sum(1 for r in rows if r["type"] == "Indicator"),
        "subrows": sum(1 for r in rows if r["type"] is None),
        "rows": rows,
    }

OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print("sheets", len(report))
for k, v in report.items():
    print(k, "M", v["metrics"], "I", v["indicators"], "sub", v["subrows"])
