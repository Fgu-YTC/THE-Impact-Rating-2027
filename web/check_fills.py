# -*- coding: utf-8 -*-
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent


def fill_info(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return "none"
    fg = f.fgColor
    if fg is None:
        return str(f.fill_type)
    if fg.type == "rgb":
        return f"{f.fill_type}:{fg.rgb}"
    if fg.type == "theme":
        return f"{f.fill_type}:theme{fg.theme}:tint{fg.tint}"
    if fg.type == "indexed":
        return f"{f.fill_type}:idx{fg.indexed}"
    return f"{f.fill_type}:{fg.type}"


def scan(path):
    wb = load_workbook(path)
    ctr = Counter()
    samples = []
    for name in wb.sheetnames:
        if not name.strip().startswith("SDG"):
            continue
        ws = wb[name]
        for row in range(3, ws.max_row + 1):
            for col in range(1, 8):
                info = fill_info(ws.cell(row, col))
                if info != "none":
                    ctr[info] += 1
                    if len(samples) < 40:
                        a = ws.cell(row, 1).value
                        c = ws.cell(row, 3).value
                        samples.append(
                            (
                                name,
                                row,
                                col,
                                a,
                                (str(c)[:50] if c else None),
                                info,
                            )
                        )
    print("FILE", path.name)
    print("fill counts", ctr.most_common(25))
    for s in samples:
        print(s)


def scan_summary(path):
    wb = load_workbook(path)
    if "summary" not in wb.sheetnames:
        return
    ws = wb["summary"]
    print("\n=== summary evidence ===")
    # group by whether Evidence Required is set
    need = 0
    no_need = 0
    headers = 0
    for row in range(2, ws.max_row + 1):
        ref = ws.cell(row, 1).value
        title = ws.cell(row, 2).value
        fmt = ws.cell(row, 3).value
        evid = ws.cell(row, 4).value
        if not title:
            continue
        # metric headers often have no format
        if fmt is None and evid is None:
            headers += 1
            continue
        if str(evid).upper() == "YES":
            need += 1
        elif str(evid).upper() == "NO":
            no_need += 1
            print("NO evidence:", ref, str(title)[:60], fmt)
        else:
            print("other:", ref, str(title)[:40], fmt, evid)
    print("headers/section", headers, "evidence YES", need, "evidence NO", no_need)


if __name__ == "__main__":
    orig = ROOT / "Impact 2027-data collection support v1.xlsx"
    scan(orig)
    scan_summary(orig)
    for p in ROOT.glob("Impact 2027*雙語*.xlsx"):
        print("\n--- bilingual ---")
        scan(p)
