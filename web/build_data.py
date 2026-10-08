# -*- coding: utf-8 -*-
"""Build questions.json and methodology-zh.json for the web app."""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WEB = Path(__file__).resolve().parent
DATA = WEB / "data"

# Resolve bilingual excel by suffix pattern
xlsx = None
for p in ROOT.glob("Impact 2027-data collection support v1*.xlsx"):
    if "雙語" in p.name or "提问" in p.name or p.stat().st_mtime:
        # prefer bilingual
        if "雙語" in p.name:
            xlsx = p
            break
if xlsx is None:
    for p in ROOT.glob("Impact 2027-data collection support v1*.xlsx"):
        if p.name != "Impact 2027-data collection support v1.xlsx":
            xlsx = p
            break
if xlsx is None:
    xlsx = ROOT / "Impact 2027-data collection support v1.xlsx"

print("Excel:", xlsx)


def sheet_sdg_num(name: str):
    m = re.match(r"SDG\s*(\d+)", name.strip(), re.I)
    return int(m.group(1)) if m else None


def build_questions():
    wb = load_workbook(xlsx, data_only=True)
    questions = []
    sdg_meta = {}

    for sheet_name in wb.sheetnames:
        sdg = sheet_sdg_num(sheet_name)
        if sdg is None:
            continue
        ws = wb[sheet_name]
        title = ws["A1"].value or f"SDG{sdg}"
        title_zh = ws["D1"].value or ""
        sdg_meta[str(sdg)] = {"en": str(title).strip(), "zh": str(title_zh).strip() if title_zh else ""}

        current_metric = None
        current_indicator = None
        q_index = 0

        for row in range(3, ws.max_row + 1):
            typ = ws.cell(row, 1).value
            ref = ws.cell(row, 2).value
            en = ws.cell(row, 3).value
            zh = ws.cell(row, 4).value
            if not en and not zh and not typ and not ref:
                continue

            typ_s = str(typ).strip() if typ else ""
            ref_s = str(ref).strip() if ref is not None else ""
            en_s = str(en).strip() if en else ""
            zh_s = str(zh).strip() if zh else ""

            if typ_s == "Metric":
                current_metric = {"ref": ref_s, "en": en_s, "zh": zh_s}
                continue

            if typ_s == "Indicator":
                q_index += 1
                current_indicator = {
                    "id": f"sdg{sdg}-{ref_s or q_index}",
                    "sdg": sdg,
                    "type": "indicator",
                    "ref": ref_s,
                    "en": en_s,
                    "zh": zh_s,
                    "metricRef": current_metric["ref"] if current_metric else "",
                    "metricEn": current_metric["en"] if current_metric else "",
                    "metricZh": current_metric["zh"] if current_metric else "",
                    "answerType": "yesno",  # default; sub-options may be picklist
                    "options": [],
                }
                questions.append(current_indicator)
                continue

            # Sub-row under indicator (option or continuous field)
            if en_s and current_indicator is not None:
                # Heuristic: Number of... -> continuous; short labels -> option
                lower = en_s.lower()
                if lower.startswith("number of") or lower.startswith("total ") or lower.startswith("amount of") or lower.startswith("volume of") or lower.startswith("university expenditure") or lower.startswith("research income") or lower.startswith("campus population") or "per sqm" in lower or lower.startswith("achieve by"):
                    current_indicator["answerType"] = "continuous"
                    current_indicator["options"].append({"en": en_s, "zh": zh_s, "kind": "value"})
                else:
                    if current_indicator["answerType"] == "yesno" and current_indicator["options"]:
                        current_indicator["answerType"] = "picklist"
                    elif current_indicator["answerType"] == "yesno" and not current_indicator["options"]:
                        # first option turns it into picklist-capable yes/no with sub choices
                        current_indicator["answerType"] = "picklist"
                    current_indicator["options"].append({"en": en_s, "zh": zh_s, "kind": "option"})

    DATA.mkdir(parents=True, exist_ok=True)
    out = {"sdgs": sdg_meta, "questions": questions}
    path = DATA / "questions.json"
    path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {path} ({len(questions)} questions)")
    return out


SDG_TITLES_ZH = {
    1: "消除貧窮",
    2: "消除飢餓",
    3: "良好健康與福祉",
    4: "優質教育",
    5: "性別平等",
    6: "潔淨水與衛生",
    7: "可負擔潔淨能源",
    8: "尊嚴勞動與經濟成長",
    9: "產業、創新與基礎設施",
    10: "減少不平等",
    11: "永續城市與社區",
    12: "負責任消費與生產",
    13: "氣候行動",
    14: "水下生命",
    15: "陸域生命",
    16: "和平、正義與健全制度",
    17: "促進目標實現的夥伴關係",
}


def to_traditional(s: str) -> str:
    from zhconv import convert

    return convert(s, "zh-tw")


def extract_pdf_sections(pdf_path: Path, start_patterns):
    """Return {sdg_num: body} from PDF using start_patterns (list of regex)."""
    text = ""
    try:
        import pdfplumber

        with pdfplumber.open(pdf_path) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        print(f"PDF {pdf_path.name} length: {len(text)}")
    except Exception as e:
        print("PDF extract failed:", e)
        return {}

    starts = []
    for pat in start_patterns:
        starts = list(re.finditer(pat, text, flags=re.I))
        if starts:
            break
    out = {}
    for i, m in enumerate(starts):
        n = int(m.group(1))
        end = starts[i + 1].start() if i + 1 < len(starts) else len(text)
        body = text[m.start() : end].strip()
        if len(body) > 4500:
            body = body[:4500].rstrip() + "…"
        key = str(n)
        if key not in out or len(body) > len(out[key]):
            out[key] = body
    return out


def build_methodology():
    """Extract SDG sections; convert Simplified → Traditional for ZH."""
    zh_pdf = ROOT / "Chinese_THE_Impact_Rankings_METHODOLOGY_2027.pdf"
    en_pdf = ROOT / "THE.SustainabilityImpactRatings.METHODOLOGY.2027.v1.2.pdf"

    zh_raw = extract_pdf_sections(
        zh_pdf,
        [
            r"我们为什么衡量\s*SDG\s*(1[0-7]|[1-9])",
            r"(?:^|\n)\s*SDG\s*(1[0-7]|[1-9])\s*(?:\n|:)",
        ],
    )
    en_raw = extract_pdf_sections(
        en_pdf,
        [
            # Chapter open: "SDG N\nTitle\nWhy we measure"
            r"(?m)^SDG\s+(1[0-7]|[1-9])\s*\n[^\n]{3,80}\nWhy we measure",
            r"Why we measure\s*SDG\s*(1[0-7]|[1-9])",
        ],
    )

    methodology_zh = {}
    methodology_en = {}
    for n in range(1, 18):
        key = str(n)
        title_zh = f"SDG{n}：{SDG_TITLES_ZH.get(n, '')}"
        body_zh = zh_raw.get(key, "")
        if len(body_zh) < 80:
            body_zh = (
                f"本區顯示 SDG{n}（{SDG_TITLES_ZH.get(n, '')}）相關方法論摘要。"
                "請依指標定義提供證據，並確認政策／證據日期符合 THE 當年度要求。"
            )
        else:
            body_zh = to_traditional(body_zh)
            # strip repeated guide footers noise lightly
            body_zh = re.sub(
                r"泰晤士高等教育世界大學可持續性影響力評級用戶指南2027\s*",
                "\n",
                body_zh,
            )
            body_zh = re.sub(r"\n{3,}", "\n\n", body_zh).strip()

        methodology_zh[key] = {"title": title_zh, "body": body_zh}

        body_en = en_raw.get(key, "")
        if len(body_en) < 80:
            body_en = (
                f"Methodology notes for SDG {n}. "
                "Please refer to THE Sustainability Impact Ratings Methodology 2027 for full definitions and evidence requirements."
            )
        methodology_en[key] = {
            "title": f"SDG {n}",
            "body": body_en,
        }

    path_zh = DATA / "methodology-zh.json"
    path_en = DATA / "methodology-en.json"
    path_zh.write_text(json.dumps(methodology_zh, ensure_ascii=False, indent=2), encoding="utf-8")
    path_en.write_text(json.dumps(methodology_en, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {path_zh} and {path_en}")

def build_units():
    units = [
        {"id": "unit-placeholder-1", "name": "（佔位）單位 A"},
        {"id": "unit-placeholder-2", "name": "（佔位）單位 B"},
        {"id": "unit-placeholder-3", "name": "（佔位）單位 C"},
    ]
    path = DATA / "units.json"
    path.write_text(json.dumps(units, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {path}")


if __name__ == "__main__":
    build_questions()
    build_methodology()
    build_units()
