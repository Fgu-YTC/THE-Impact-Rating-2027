# -*- coding: utf-8 -*-
"""
Rebuild questions.json from ORIGINAL Excel (source of truth for structure).
Chinese text is matched by indicator/metric ref from existing bilingual data when possible.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from zhconv import convert

ROOT = Path(__file__).resolve().parent.parent
DATA = Path(__file__).resolve().parent / "data"
ORIG = ROOT / "Impact 2027-data collection support v1.xlsx"

SDG_TITLES_ZH = {
    1: "消除貧窮",
    2: "消除飢餓",
    3: "良好健康與福祉",
    4: "優質教育",
    5: "性別平等",
    6: "潔淨水與衛生",
    7: "可負擔潔淨能源",
    8: "尊嚴勞動與經濟成長",
    9: "產業、創新與基礎設施",
    10: "減少不平等",
    11: "永續城市與社區",
    12: "負責任消費與生產",
    13: "氣候行動",
    14: "水下生命",
    15: "陸域生命",
    16: "和平、正義與健全制度",
    17: "促進目標實現的夥伴關係",
}


def sheet_sdg(name: str):
    m = re.match(r"SDG\s*(\d+)", name.strip(), re.I)
    return int(m.group(1)) if m else None


def clean(s):
    if s is None:
        return ""
    s = str(s).replace("_x000D_", "").replace("\r\n", "\n").replace("\r", "\n")
    for a, b in [
        ("\u2018", "'"),
        ("\u2019", "'"),
        ("\u201c", '"'),
        ("\u201d", '"'),
        ("\u2013", "-"),
        ("\u2014", "-"),
        ("\uf0a7", "•"),
        ("\u00a0", " "),
    ]:
        s = s.replace(a, b)
    return "\n".join(line.rstrip() for line in s.split("\n")).strip()


def load_zh_lookup():
    """Map ref -> {en, zh} and also metricRef -> zh from previous questions.json if present."""
    path = DATA / "questions.json"
    by_ref = {}
    by_en = {}
    if not path.exists():
        return by_ref, by_en
    data = json.loads(path.read_text(encoding="utf-8"))
    for q in data.get("questions", []):
        ref = str(q.get("ref") or "").strip()
        if ref:
            by_ref[ref] = q
        en = clean(q.get("en"))
        if en:
            by_en[en] = q
        # options
        for opt in q.get("options") or []:
            oen = clean(opt.get("en"))
            if oen and opt.get("zh"):
                by_en[oen] = {"zh": opt["zh"], "en": oen}
    # metric titles stored on questions
    for q in data.get("questions", []):
        mref = str(q.get("metricRef") or "").strip()
        if mref and q.get("metricZh"):
            by_ref.setdefault(
                mref,
                {"ref": mref, "en": q.get("metricEn"), "zh": q.get("metricZh")},
            )
    return by_ref, by_en


def zh_for(en: str, ref: str | None, by_ref, by_en) -> str:
    en_c = clean(en)
    if ref and ref in by_ref and by_ref[ref].get("zh"):
        return clean(by_ref[ref]["zh"])
    if en_c in by_en and by_en[en_c].get("zh"):
        return clean(by_en[en_c]["zh"])
    # first-line match
    first = en_c.split("\n")[0].strip()
    for k, v in by_en.items():
        if k.split("\n")[0].strip() == first and v.get("zh"):
            return clean(v["zh"])
    return ""


def infer_answer_type(label: str, options: list) -> str:
    if not options:
        return "yesno"
    kinds = {o.get("kind") for o in options}
    if "value" in kinds:
        return "continuous"
    return "picklist"


def classify_subrow(label: str) -> str:
    lower = label.lower()
    value_starts = (
        "number of",
        "total ",
        "amount of",
        "volume of",
        "university expenditure",
        "research income",
        "campus population",
        "university floor",
        "achieve by",
    )
    if any(lower.startswith(s) for s in value_starts) or "per sqm" in lower:
        return "value"
    return "option"


def build():
    by_ref, by_en = load_zh_lookup()
    wb = load_workbook(ORIG, data_only=True)
    questions = []
    metrics = []  # non-fillable headers for UI
    sdg_meta = {}

    for sheet_name in wb.sheetnames:
        sdg = sheet_sdg(sheet_name)
        if sdg is None:
            continue
        ws = wb[sheet_name]
        title = clean(ws["A1"].value) or f"SDG{sdg}"
        sdg_meta[str(sdg)] = {
            "en": title,
            "zh": f"SDG{sdg}：{SDG_TITLES_ZH.get(sdg, '')}",
        }

        current_metric = None
        current_indicator = None

        for row in range(3, ws.max_row + 1):
            typ = ws.cell(row, 1).value
            ref = ws.cell(row, 2).value
            en = clean(ws.cell(row, 3).value)
            typ_s = str(typ).strip() if typ else ""
            ref_s = str(ref).strip() if ref is not None else ""

            if not typ_s and not ref_s and not en:
                continue

            if typ_s == "Metric":
                current_metric = {
                    "id": f"metric-sdg{sdg}-{ref_s}",
                    "sdg": sdg,
                    "type": "metric",
                    "fillable": False,  # 灰底：不用填
                    "ref": ref_s,
                    "en": en,
                    "zh": zh_for(en, ref_s, by_ref, by_en),
                }
                metrics.append(current_metric)
                current_indicator = None
                continue

            if typ_s == "Indicator":
                q = {
                    "id": f"sdg{sdg}-{ref_s}",
                    "sdg": sdg,
                    "type": "indicator",
                    "fillable": True,  # 白底：要填
                    "ref": ref_s,
                    "en": en,
                    "zh": zh_for(en, ref_s, by_ref, by_en),
                    "metricRef": current_metric["ref"] if current_metric else "",
                    "metricEn": current_metric["en"] if current_metric else "",
                    "metricZh": current_metric["zh"] if current_metric else "",
                    "answerType": "yesno",
                    "options": [],
                }
                questions.append(q)
                current_indicator = q
                continue

            # sub-row under current indicator
            if en and current_indicator is not None:
                kind = classify_subrow(en)
                opt = {
                    "en": en,
                    "zh": zh_for(en, None, by_ref, by_en),
                    "kind": kind,
                }
                current_indicator["options"].append(opt)
                current_indicator["answerType"] = infer_answer_type(
                    current_indicator["en"], current_indicator["options"]
                )

    # Fill missing metricZh/zh from first line heuristics already done
    missing_zh = [q["ref"] for q in questions if not q.get("zh")]
    missing_metric = [m["ref"] for m in metrics if not m.get("zh")]

    out = {
        "source": ORIG.name,
        "note": "Structure from original Excel. Metric fillable=false (灰底不用填); Indicator fillable=true (白底要填).",
        "sdgs": sdg_meta,
        "metrics": metrics,
        "questions": questions,
        "counts": {
            "metrics": len(metrics),
            "indicators": len(questions),
            "missingQuestionZh": len(missing_zh),
            "missingMetricZh": len(missing_metric),
        },
    }
    path = DATA / "questions.json"
    path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Wrote", path)
    print(json.dumps(out["counts"], ensure_ascii=False, indent=2))
    if missing_zh[:15]:
        print("sample missing zh refs:", missing_zh[:15])


if __name__ == "__main__":
    build()
