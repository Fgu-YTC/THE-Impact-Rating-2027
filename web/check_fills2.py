# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
wb = load_workbook(ROOT / "Impact 2027-data collection support v1.xlsx")


def fill_info(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return "none"
    fg = f.fgColor
    if fg is None:
        return str(f.fill_type)
    if fg.type == "rgb":
        return f"{f.fill_type}:{fg.rgb}"
    if fg.type == "theme":
        return f"{f.fill_type}:theme{fg.theme}:tint{round(fg.tint or 0, 3)}"
    if fg.type == "indexed":
        return f"{f.fill_type}:idx{fg.indexed}"
    return f"{f.fill_type}:{fg.type}"


# Compare Metric vs Indicator row fills across columns for a few sheets
for name in ["SDG1 ", "SDG5", "SDG9", "SDG17"]:
    ws = wb[name]
    print("\n====", name, "====")
    for row in range(2, min(ws.max_row + 1, 25)):
        typ = ws.cell(row, 1).value
        ref = ws.cell(row, 2).value
        label = ws.cell(row, 3).value
        fills = [fill_info(ws.cell(row, c)) for c in range(1, 8)]
        uniq = set(fills)
        print(
            f"R{row} type={typ!r} ref={ref!r} fills={fills} label={(str(label)[:40] if label else None)!r}"
        )

# example sheet
print("\n==== example ====")
ws = wb["example"]
for row in range(1, ws.max_row + 1):
    typ = ws.cell(row, 1).value
    fills = [fill_info(ws.cell(row, c)) for c in range(1, 8)]
    label = ws.cell(row, 3).value
    print(f"R{row} type={typ!r} fills={fills} {(str(label)[:45] if label else None)!r}")
