# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(
    r"E:\奕廷\永續辦業務\大學排名\THE\2027\Impact Rating 2027\Impact 2027-data collection support v1_雙語提問.xlsx"
)
lines = []
ws = wb["SDG1 "]
lines.append("=== SDG1 ===")
lines.append(f"A1={ws['A1'].value}")
lines.append(f"D1={ws['D1'].value}")
lines.append(f"headers={[ws.cell(2, c).value for c in range(1, 9)]}")
for r in range(3, 12):
    lines.append(f"R{r}")
    lines.append(f"  EN: {ws.cell(r, 3).value}")
    lines.append(f"  ZH: {ws.cell(r, 4).value}")

ws = wb["SDG17"]
lines.append("=== SDG17 R4 ===")
lines.append(f"EN: {ws.cell(4, 3).value}")
lines.append(f"ZH: {ws.cell(4, 4).value}")

empty = 0
total = 0
for name in wb.sheetnames:
    if not name.strip().startswith("SDG"):
        continue
    ws = wb[name]
    for r in range(3, ws.max_row + 1):
        en = ws.cell(r, 3).value
        zh = ws.cell(r, 4).value
        if en:
            total += 1
            if not zh:
                empty += 1
                lines.append(f"MISSING {name} R{r}: {en!r}")

lines.append(f"SDG EN cells={total} missing ZH={empty}")
path = r"C:\Users\fgu\.cursor\projects\e-THE-2027-Impact-Rating-2027\agent-tools\verify_utf8.txt"
open(path, "w", encoding="utf-8").write("\n".join(lines))
print("total", total, "empty", empty)
print("wrote", path)
